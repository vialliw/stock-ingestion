"""Tests for the standalone APScheduler-based daily load scheduler."""

from __future__ import annotations

import io
import json
import os
import signal
import subprocess
import sys
from pathlib import Path
from unittest.mock import ANY, MagicMock, patch

import duckdb
import openpyxl
import pyarrow as pa
import pyarrow.flight as flight
import pytest

# Add scripts dir to path so daily_load_once can be imported (same
# pattern as tests/test_poll_tasks.py).
SCRIPTS_DIR = Path(__file__).parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import daily_load_once
import report_top25_active

# run_flight_server's module header chdirs to the repo root (it is meant
# to be launched by Task Scheduler from any cwd); restore the test
# session's cwd after importing so collection stays hermetic.
_cwd_before_rfs_import = os.getcwd()
import run_flight_server
os.chdir(_cwd_before_rfs_import)

from src.config import PipelineConfig, SchedulerConfig
from src.scheduler import PipelineScheduler


@pytest.fixture
def enabled_config(tmp_path):
    """A config with the scheduler enabled and an isolated database path."""
    return PipelineConfig(
        database={"duckdb_dir": str(tmp_path), "duckdb_filename": "test.duckdb"},
        scheduler=SchedulerConfig(enabled=True),
    )


@pytest.fixture
def disabled_config(tmp_path):
    """A config with the scheduler disabled."""
    return PipelineConfig(
        database={"duckdb_dir": str(tmp_path), "duckdb_filename": "test.duckdb"},
        scheduler=SchedulerConfig(enabled=False),
    )


def test_parse_schedule_time():
    assert PipelineScheduler._parse_schedule_time("09:05") == (9, 5)
    assert PipelineScheduler._parse_schedule_time("23:59") == (23, 59)
    assert PipelineScheduler._parse_schedule_time("garbage") == (16, 30)


def test_scheduler_disabled_does_not_start(disabled_config):
    scheduler = PipelineScheduler(disabled_config)
    with patch("src.scheduler.BackgroundScheduler") as mock_cls:
        scheduler.start(block=False)
        mock_cls.assert_not_called()


def test_scheduler_start_adds_cron_job(enabled_config):
    scheduler = PipelineScheduler(enabled_config)
    mock_scheduler = MagicMock()
    with patch("src.scheduler.BackgroundScheduler", return_value=mock_scheduler):
        scheduler.start(block=False)

    mock_scheduler.add_job.assert_called_once()
    args, kwargs = mock_scheduler.add_job.call_args
    assert args[1] == "cron"
    assert kwargs["hour"] == 16
    assert kwargs["minute"] == 30
    assert kwargs["timezone"] == "America/New_York"
    assert kwargs["id"] == "daily_load"
    assert kwargs["replace_existing"] is True
    mock_scheduler.start.assert_called_once()

    scheduler.shutdown()


def test_scheduler_start_blocks_until_shutdown(enabled_config):
    scheduler = PipelineScheduler(enabled_config)
    scheduler._scheduler = MagicMock()
    # Unblock the wait loop after the first iteration.
    scheduler._shutdown_event.set()
    scheduler.start(block=True)


@patch("src.scheduler.mcal.get_calendar")
@patch("src.scheduler.DatabaseManager")
@patch("src.scheduler.IngestionEngine")
@patch("src.scheduler.time.sleep")
def test_run_once_market_closed_skips(
    mock_sleep, mock_engine_cls, mock_db_cls, mock_get_calendar, enabled_config
):
    mock_cal = MagicMock()
    mock_cal.schedule.return_value.empty = True
    mock_get_calendar.return_value = mock_cal

    scheduler = PipelineScheduler(enabled_config)
    scheduler.run_once()

    mock_get_calendar.assert_called_once()
    mock_db_cls.assert_called_once_with(enabled_config.db_path)
    mock_engine_cls.assert_not_called()
    mock_db_cls.return_value.close.assert_called_once()


@patch("src.scheduler.mcal.get_calendar")
@patch("src.scheduler.DatabaseManager")
@patch("src.scheduler.IngestionEngine")
@patch("src.scheduler.time.sleep")
def test_run_once_market_open_runs_daily_load(
    mock_sleep, mock_engine_cls, mock_db_cls, mock_get_calendar, enabled_config
):
    mock_cal = MagicMock()
    mock_cal.schedule.return_value.empty = False
    mock_get_calendar.return_value = mock_cal

    mock_engine = MagicMock()
    mock_engine.run_daily_load.return_value = {
        "status": "SUCCESS",
        "rows_upserted": 42,
    }
    mock_engine_cls.return_value = mock_engine

    scheduler = PipelineScheduler(enabled_config)
    scheduler.run_once()

    mock_engine_cls.assert_called_once_with(mock_db_cls.return_value, enabled_config)
    mock_engine.run_daily_load.assert_called_once()
    mock_db_cls.return_value.close.assert_called_once()


@patch("src.scheduler.mcal.get_calendar")
@patch("src.scheduler.DatabaseManager")
@patch("src.scheduler.IngestionEngine")
@patch("src.scheduler.time.sleep")
def test_run_once_retries_then_succeeds_on_db_lock(
    mock_sleep, mock_engine_cls, mock_db_cls, mock_get_calendar, enabled_config
):
    mock_cal = MagicMock()
    mock_cal.schedule.return_value.empty = False
    mock_get_calendar.return_value = mock_cal

    mock_db = MagicMock()
    mock_db_cls.side_effect = [
        duckdb.IOException("locked"),
        duckdb.IOException("locked"),
        mock_db,
    ]

    scheduler = PipelineScheduler(enabled_config)
    scheduler.run_once()

    assert mock_db_cls.call_count == 3
    assert mock_sleep.call_count == 2
    mock_engine_cls.assert_called_once_with(mock_db, enabled_config)
    mock_db.close.assert_called_once()


@patch("src.scheduler.mcal.get_calendar")
@patch("src.scheduler.DatabaseManager")
@patch("src.scheduler.IngestionEngine")
@patch("src.scheduler.time.sleep")
def test_run_once_retries_exhausted_logs_error(
    mock_sleep, mock_engine_cls, mock_db_cls, mock_get_calendar, enabled_config
):
    mock_cal = MagicMock()
    mock_cal.schedule.return_value.empty = False
    mock_get_calendar.return_value = mock_cal

    mock_db_cls.side_effect = duckdb.IOException("locked")

    scheduler = PipelineScheduler(enabled_config)
    scheduler.run_once()

    assert mock_db_cls.call_count == 4
    assert mock_sleep.call_count == 3
    mock_engine_cls.assert_not_called()


@patch("src.scheduler.signal.signal")
def test_install_signal_handlers(mock_signal, enabled_config):
    scheduler = PipelineScheduler(enabled_config)
    scheduler.install_signal_handlers()
    mock_signal.assert_any_call(signal.SIGINT, ANY)
    mock_signal.assert_any_call(signal.SIGTERM, ANY)


@patch("src.scheduler.IngestionEngine")
def test_execute_daily_load_force_bypasses_market_gate(
    mock_engine_cls, enabled_config
):
    """F7-T15: force=True bypasses the NYSE market-closed gate entirely
    and refreshes the configured ``ingestion.daily_lookback_days`` window
    so a missed trading day is caught up (local one-shot runner path)."""
    mock_engine = MagicMock()
    mock_engine.run_daily_load.return_value = {
        "status": "SUCCESS",
        "rows_upserted": 7,
    }
    mock_engine_cls.return_value = mock_engine

    scheduler = PipelineScheduler(enabled_config)
    db_manager = MagicMock()
    with patch.object(
        PipelineScheduler, "_is_market_open_today", return_value=False
    ) as mock_gate:
        scheduler._execute_daily_load(db_manager, force=True)

    # The market-open gate was never consulted.
    mock_gate.assert_not_called()
    mock_engine_cls.assert_called_once_with(db_manager, enabled_config)
    mock_engine.run_daily_load.assert_called_once_with(
        lookback_days=enabled_config.ingestion.daily_lookback_days,
        force=True,
    )


@patch("src.scheduler.mcal.get_calendar")
@patch("src.scheduler.IngestionEngine")
def test_execute_daily_load_default_keeps_gate_and_lookback_1(
    mock_engine_cls, mock_get_calendar, enabled_config
):
    """Daemon default (force=False) is unchanged: the market-open gate is
    consulted and the engine is called with lookback_days=1 and no force
    flag, so the APScheduler path still skips on market-closed days."""
    mock_cal = MagicMock()
    mock_cal.schedule.return_value.empty = False
    mock_get_calendar.return_value = mock_cal

    mock_engine = MagicMock()
    mock_engine.run_daily_load.return_value = {
        "status": "SUCCESS",
        "rows_upserted": 1,
    }
    mock_engine_cls.return_value = mock_engine

    scheduler = PipelineScheduler(enabled_config)
    scheduler._execute_daily_load(MagicMock())

    mock_get_calendar.assert_called_once()
    mock_engine.run_daily_load.assert_called_once_with(lookback_days=1)


# =====================================================================
# F7-T16: scripts/daily_load_once.py — service-first Flight routing
# =====================================================================


def _mock_flight_module():
    """A stand-in for the ``pyarrow.flight`` module inside
    ``daily_load_once``: real value/exception classes (needed for
    ``except flight.<X>Error`` clauses and Action construction) but a
    mocked ``FlightClient``.
    """
    mock_flight = MagicMock()
    mock_flight.Action = flight.Action
    mock_flight.FlightCallOptions = flight.FlightCallOptions
    mock_flight.FlightError = flight.FlightError
    mock_flight.FlightServerError = flight.FlightServerError
    mock_flight.FlightTimedOutError = flight.FlightTimedOutError
    mock_flight.FlightUnavailableError = flight.FlightUnavailableError
    return mock_flight


def _flight_result(payload: dict):
    """Build a mock flight.Result-like object carrying a JSON body."""
    result = MagicMock()
    result.body.to_pybytes.return_value = json.dumps(payload).encode("utf-8")
    return result


class TestDailyLoadOnceFlightRouting:
    """daily_load_once._run_once routes the daily load through the
    running Flight service when reachable, and falls back to the direct
    DuckDB path only when the service is unreachable."""

    @pytest.fixture
    def config(self, tmp_path):
        """Enabled scheduler config with an isolated database path."""
        return PipelineConfig(
            database={
                "duckdb_dir": str(tmp_path),
                "duckdb_filename": "test.duckdb",
            },
            scheduler=SchedulerConfig(enabled=True),
        )

    @pytest.fixture(autouse=True)
    def _stub_chained_report(self):
        """F7-T20: ``_run_once`` chains ``report_top25_active`` after a
        successful load.  Stub it out so these routing tests stay
        hermetic (no real Flight query / no stray xlsx under cwd);
        chaining itself is covered by ``TestDailyLoadOnceReportChaining``.
        """
        with patch.object(
            daily_load_once, "report_top25_active"
        ) as mock_report:
            mock_report.run_report.return_value = 0
            yield mock_report

    @pytest.mark.parametrize("status", ["SUCCESS", "PARTIAL"])
    def test_flight_success_no_direct_db_open(self, config, status):
        """Service up + SUCCESS/PARTIAL -> exit 0 without ever opening
        DuckDB directly (the service stays the single writer)."""
        mock_flight = _mock_flight_module()
        mock_client = mock_flight.FlightClient.return_value
        mock_client.do_action.return_value = iter(
            [_flight_result({"status": status, "rows_upserted": 5})]
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ) as mock_sched_cls:
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_sched = mock_sched_cls.return_value
        mock_sched._create_db_manager.assert_not_called()
        mock_sched._execute_daily_load.assert_not_called()

        # The action sent to the service is daily_load with force=true.
        action_arg = mock_client.do_action.call_args.args[0]
        assert action_arg.type == "daily_load"
        assert json.loads(action_arg.body.to_pybytes()) == {"force": True}

    @pytest.mark.parametrize(
        "error",
        [
            flight.FlightServerError("server exploded"),
            flight.FlightTimedOutError("deadline exceeded"),
        ],
        ids=["server_error", "timeout"],
    )
    def test_flight_server_error_or_timeout_returns_1(self, config, error):
        """Service up but the action fails -> exit 1, no direct-DB
        fallback (the service handled the request)."""
        mock_flight = _mock_flight_module()
        mock_client = mock_flight.FlightClient.return_value
        mock_client.do_action.side_effect = error

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ) as mock_sched_cls:
            code = daily_load_once._run_once(config)

        assert code == 1
        mock_sched_cls.return_value._create_db_manager.assert_not_called()

    def test_flight_non_success_status_returns_1(self, config):
        """A reachable service that answers with a non-success status
        (e.g. SKIPPED from a stale server without force support) ->
        exit 1, no direct-DB fallback."""
        mock_flight = _mock_flight_module()
        mock_client = mock_flight.FlightClient.return_value
        mock_client.do_action.return_value = iter(
            [_flight_result({"status": "SKIPPED", "reason": "Market closed"})]
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ) as mock_sched_cls:
            code = daily_load_once._run_once(config)

        assert code == 1
        mock_sched_cls.return_value._create_db_manager.assert_not_called()

    def test_flight_unreachable_falls_back_to_direct(self, config):
        """Connection failure (server down) -> log + fall back to the
        existing direct path _execute_daily_load(db_manager, force=True).
        """
        mock_flight = _mock_flight_module()
        mock_client = mock_flight.FlightClient.return_value
        mock_client.wait_for_available.side_effect = (
            flight.FlightUnavailableError("connection refused")
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ) as mock_sched_cls:
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_sched = mock_sched_cls.return_value
        mock_sched._create_db_manager.assert_called_once()
        mock_sched._execute_daily_load.assert_called_once_with(
            mock_sched._create_db_manager.return_value, force=True
        )
        mock_sched._create_db_manager.return_value.close.assert_called_once()

    def test_flight_unavailable_mid_call_falls_back(self, config):
        """If the server drops between the availability probe and the
        action call, the direct path still takes over."""
        mock_flight = _mock_flight_module()
        mock_client = mock_flight.FlightClient.return_value
        mock_client.do_action.side_effect = flight.FlightUnavailableError(
            "server went away"
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ) as mock_sched_cls:
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_sched = mock_sched_cls.return_value
        mock_sched._execute_daily_load.assert_called_once_with(
            mock_sched._create_db_manager.return_value, force=True
        )

    def test_scheduler_disabled_returns_0_without_flight(self, tmp_path):
        """enabled=false -> 0 and the Flight service is never consulted."""
        config = PipelineConfig(
            database={
                "duckdb_dir": str(tmp_path),
                "duckdb_filename": "test.duckdb",
            },
            scheduler=SchedulerConfig(enabled=False),
        )
        mock_flight = _mock_flight_module()

        with patch.object(daily_load_once, "flight", mock_flight):
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_flight.FlightClient.assert_not_called()


# =====================================================================
# F7-T19: run_flight_server._is_db_lock_error markers + challenger parity
# =====================================================================


class TestIsDbLockError:
    """Unit tests for ``scripts/run_flight_server._is_db_lock_error``.

    Challenger finding N2 (F7-T18 review): the previous bare ``'lock'``
    marker substring-matched the quoted DB path inside the IOException
    message, so a non-contention error under a directory named e.g.
    ``blocked_dir`` was misclassified as write-lock contention and
    retried for the full 600 s window.  Markers are now phrases that
    only appear in the *reason* text, and double-quoted spans (the path)
    are stripped before matching.
    """

    # Real error text captured live on this machine: duckdb 1.4.3, two
    # processes competing for the same file, Windows.
    REAL_WINDOWS_LOCK_MSG = (
        'IO Error: Cannot open file "c:\\users\\user\\appdata\\local\\temp\\'
        'tmpxyz\\test.duckdb": The process cannot access the file because '
        'it is being used by another process.\r\n\nFile is already open in \n'
        "C:\\Intel\\python\\3.12\\python.exe (PID 21204)"
    )

    # POSIX contention shape from duckdb's file-lock path (same release):
    # 'Could not set lock on file "<path>": Conflicting lock is held ...'
    POSIX_LOCK_MSG = (
        'IO Error: Could not set lock on file "/data/market.duckdb": '
        "Conflicting lock is held in /usr/bin/python3.12 (PID 4321)"
    )

    def test_real_windows_lock_message_is_lock(self):
        exc = duckdb.IOException(self.REAL_WINDOWS_LOCK_MSG)
        assert run_flight_server._is_db_lock_error(exc) is True

    def test_posix_lock_message_is_lock(self):
        exc = duckdb.IOException(self.POSIX_LOCK_MSG)
        assert run_flight_server._is_db_lock_error(exc) is True

    def test_value_error_is_never_lock(self):
        # isinstance gate: a non-IOException carrying lock wording -> False
        exc = ValueError(
            "Could not set lock on file: being used by another process"
        )
        assert run_flight_server._is_db_lock_error(exc) is False

    def test_non_lock_ioexception_is_not_lock(self):
        exc = duckdb.IOException(
            'IO Error: Cannot open file "c:\\data\\m.duckdb": '
            "The system cannot find the path specified."
        )
        assert run_flight_server._is_db_lock_error(exc) is False

    def test_missing_dir_with_lock_in_path_is_not_lock(self, tmp_path):
        """N2 regression: a real duckdb IOException whose *path* contains
        'lock' ('blocked_dir_missing') must NOT classify as contention."""
        missing = tmp_path / "blocked_dir_missing" / "x.duckdb"
        with pytest.raises(duckdb.IOException) as excinfo:
            duckdb.connect(str(missing))
        # Sanity: the message really does contain 'lock' via the path.
        assert "lock" in str(excinfo.value).lower()
        assert run_flight_server._is_db_lock_error(excinfo.value) is False

    def test_real_two_process_lock_under_lock_named_dir(self, tmp_path):
        """Live check: a real cross-process write lock on a DB under a
        directory named 'blocked_dir' is still classified as contention —
        the marker must come from the reason text, never the path."""
        db_dir = tmp_path / "blocked_dir"
        db_dir.mkdir()
        db_path = str(db_dir / "m.duckdb")
        holder = subprocess.Popen(
            [
                sys.executable,
                "-c",
                "import duckdb, sys, time; "
                "c = duckdb.connect(sys.argv[1]); "
                "c.execute('CREATE TABLE t(i int)'); "
                "print('ready', flush=True); time.sleep(30)",
                db_path,
            ],
            stdout=subprocess.PIPE,
            text=True,
        )
        try:
            assert holder.stdout.readline().strip() == "ready"
            with pytest.raises(duckdb.IOException) as excinfo:
                duckdb.connect(db_path)
            # The quoted path really does contain 'lock' ('blocked_dir').
            assert "lock" in str(excinfo.value).lower()
            assert run_flight_server._is_db_lock_error(excinfo.value) is True
        finally:
            holder.terminate()
            holder.wait(timeout=10)


class TestChallengerDailyLookbackParity:
    """F7-T19/N1: ``config_challenger.IngestionConfig.daily_lookback_days``
    parity with ``src/config.py`` (default 4, int >= 1), and the
    challenger flight server forwards the *configured* value to the
    engine — the pairing that previously raised AttributeError."""

    def test_default_is_four(self):
        from src.config_challenger import PipelineConfig as ChallengerConfig

        assert ChallengerConfig().ingestion.daily_lookback_days == 4

    def test_yaml_value_is_read(self, tmp_path):
        from src.config_challenger import PipelineConfig as ChallengerConfig

        cfg_file = tmp_path / "cfg.yaml"
        cfg_file.write_text("ingestion:\n  daily_lookback_days: 9\n")
        cfg = ChallengerConfig.from_yaml(str(cfg_file))
        assert cfg.ingestion.daily_lookback_days == 9

    @pytest.mark.parametrize("bad", ["0", "-3"])
    def test_yaml_rejects_lookback_below_one(self, tmp_path, bad):
        from src.config_challenger import PipelineConfig as ChallengerConfig

        cfg_file = tmp_path / "cfg.yaml"
        cfg_file.write_text(f"ingestion:\n  daily_lookback_days: {bad}\n")
        with pytest.raises(ValueError):
            ChallengerConfig.from_yaml(str(cfg_file))

    def test_challenger_daily_load_forwards_configured_lookback(self):
        """Challenger flight server + challenger config pairing: the
        engine must receive ``cfg.ingestion.daily_lookback_days`` — a
        non-default sentinel proves it is not a hardcoded literal, and a
        real challenger config proves the AttributeError (N1) is gone."""
        import src.flight_server_challenger as challenger_module
        from src.config_challenger import PipelineConfig as ChallengerConfig

        cfg = ChallengerConfig()
        cfg.ingestion.daily_lookback_days = 9  # non-default sentinel
        ingest_engine = MagicMock()
        ingest_engine.run_daily_load.return_value = {
            "rows_upserted": 1,
            "tickers_succeeded": ["AAPL"],
            "tickers_failed": [],
            "errors": {},
            "duration_ms": 5,
        }

        class _TestServer(challenger_module.MarketDataFlightServer):
            """Skips the port-binding base ``__init__`` (same pattern as
            ``tests/test_flight_server.py::make_server``)."""

            def __init__(self):
                self.db = MagicMock()
                self.engine = ingest_engine
                self.cfg = cfg

        server = _TestServer()
        action = MagicMock()
        action.type = "daily_load"
        action.body = None
        data = json.loads(
            server.do_action(MagicMock(), action)
            .body.to_pybytes()
            .decode("utf-8")
        )

        assert data["status"] == "SUCCESS"
        ingest_engine.run_daily_load.assert_called_once_with(
            lookback_days=9, force=False
        )


# =====================================================================
# F8-T2: run_flight_server._ensure_stdio_fallbacks — headless stdio guard
# =====================================================================


class TestEnsureStdioFallbacks:
    """Unit tests for ``scripts/run_flight_server._ensure_stdio_fallbacks``.

    The ``StockIngestion-FlightServer`` scheduled task runs with no
    console, so Python starts with ``sys.stdout is sys.stderr is None``.
    The live-observed crash was yfinance's tqdm progress-bar finalizer
    calling ``sys.stderr.flush()`` *after* a download completed —
    ``AttributeError: 'NoneType' object has no attribute 'flush'`` —
    killing a finished 10-year initial_load before upsert.  The guard
    redirects missing streams to the null device; existing streams must
    be left untouched.
    """

    def test_none_stderr_replaced_with_writable_devnull(self, monkeypatch):
        monkeypatch.setattr(sys, "stderr", None)
        run_flight_server._ensure_stdio_fallbacks()
        assert sys.stderr is not None
        # The exact call that crashed live — must not raise.
        sys.stderr.flush()
        sys.stderr.write("headless write\n")
        assert sys.stderr.name == os.devnull

    def test_none_stdout_replaced_with_writable_devnull(self, monkeypatch):
        monkeypatch.setattr(sys, "stdout", None)
        run_flight_server._ensure_stdio_fallbacks()
        assert sys.stdout is not None
        sys.stdout.flush()
        sys.stdout.write("headless write\n")
        assert sys.stdout.name == os.devnull

    def test_both_streams_none_replaced(self, monkeypatch):
        monkeypatch.setattr(sys, "stdout", None)
        monkeypatch.setattr(sys, "stderr", None)
        run_flight_server._ensure_stdio_fallbacks()
        assert sys.stdout is not None
        assert sys.stderr is not None
        sys.stdout.flush()
        sys.stderr.flush()

    def test_existing_streams_left_untouched(self, monkeypatch):
        """Idempotent no-op: real streams (console, pytest capture) must
        never be swapped out."""
        sentinel_out = io.StringIO()
        sentinel_err = io.StringIO()
        monkeypatch.setattr(sys, "stdout", sentinel_out)
        monkeypatch.setattr(sys, "stderr", sentinel_err)
        run_flight_server._ensure_stdio_fallbacks()
        assert sys.stdout is sentinel_out
        assert sys.stderr is sentinel_err

    def test_partial_none_only_missing_stream_replaced(self, monkeypatch):
        sentinel_out = io.StringIO()
        monkeypatch.setattr(sys, "stdout", sentinel_out)
        monkeypatch.setattr(sys, "stderr", None)
        run_flight_server._ensure_stdio_fallbacks()
        assert sys.stdout is sentinel_out
        assert sys.stderr is not None
        sys.stderr.flush()


# =====================================================================
# F7-T20: scripts/report_top25_active.py — top-25 active Excel report
# =====================================================================


def _report_config(tmp_path) -> PipelineConfig:
    """Config whose DB and report destination live under ``tmp_path``.

    The report is written to ``dirname(logging.log_file)`` — here
    ``<tmp>/logs/`` — so nothing touches the real C:\\Intel\\parquet.
    """
    return PipelineConfig(
        database={
            "duckdb_dir": str(tmp_path),
            "duckdb_filename": "test.duckdb",
        },
        logging={"log_file": str(tmp_path / "logs" / "run.log")},
    )


def _seed_daily_close(db_path: str, rows: list) -> None:
    """Create a fresh DuckDB file with a ``daily_close`` table holding
    ``rows`` of (date_str, ticker, open, high, low, close, volume)."""
    conn = duckdb.connect(db_path)
    try:
        conn.execute(
            """
            CREATE TABLE daily_close (
                Date DATE NOT NULL,
                Ticker VARCHAR NOT NULL,
                Open DOUBLE,
                High DOUBLE,
                Low DOUBLE,
                "Close" DOUBLE,
                Volume DOUBLE,
                PRIMARY KEY (Date, Ticker)
            )
            """
        )
        conn.executemany(
            "INSERT INTO daily_close VALUES (?, ?, ?, ?, ?, ?, ?)", rows
        )
    finally:
        conn.close()


def _mock_flight_for_report(table=None):
    """Stand-in for the ``pyarrow.flight`` module inside
    ``report_top25_active`` (same pattern as ``_mock_flight_module``)."""
    mock_flight = MagicMock()
    mock_flight.Ticket = flight.Ticket
    mock_flight.FlightCallOptions = flight.FlightCallOptions
    mock_flight.FlightError = flight.FlightError
    mock_flight.FlightServerError = flight.FlightServerError
    mock_flight.FlightTimedOutError = flight.FlightTimedOutError
    mock_flight.FlightUnavailableError = flight.FlightUnavailableError
    if table is not None:
        reader = MagicMock()
        reader.read_all.return_value = table
        mock_flight.FlightClient.return_value.do_get.return_value = reader
    return mock_flight


class TestReportTop25Active:
    """Unit tests for ``scripts/report_top25_active.py``."""

    def test_sql_has_limit_25_and_friendly_headers(self):
        """The PM-supplied query was adapted: TOP 25 via LIMIT 25,
        ordered by avg_trade_money DESC, with friendly output headers."""
        sql = report_top25_active.TOP25_SQL
        assert "LIMIT 25" in sql
        assert "ORDER BY a.avg_trade_money DESC" in sql
        assert 'AS "Ticker"' in sql
        assert 'AS "Avg Trade Money"' in sql
        assert 'AS "% of Total"' in sql
        # Window is the last 30 trading periods per ticker.
        assert "period_rank <= 30" in sql
        assert "FROM daily_close" in sql

    def test_xlsx_written_via_flight(self, tmp_path):
        """Happy path: Flight do_get -> xlsx with expected headers,
        rows, and auto-adjusted column widths (read back via openpyxl)."""
        config = _report_config(tmp_path)
        table = pa.table(
            {
                "Ticker": ["MSFT", "AAPL"],
                "Avg Trade Money": [1234.5, 678.9],
                "% of Total": [64.5, 35.5],
            }
        )
        mock_flight = _mock_flight_for_report(table)
        mock_client = mock_flight.FlightClient.return_value

        with patch.object(report_top25_active, "flight", mock_flight):
            code = report_top25_active.run_report(config)

        assert code == 0
        mock_client.wait_for_available.assert_called_once()
        mock_client.do_get.assert_called_once()
        ticket = mock_client.do_get.call_args.args[0]
        payload = json.loads(ticket.ticket.decode("utf-8"))
        assert payload["sql"] == report_top25_active.TOP25_SQL

        out_path = tmp_path / "logs" / "Report_Top_25_Active.xlsx"
        assert out_path.exists()
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert headers == ["Ticker", "Avg Trade Money", "% of Total"]
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 2
        assert data_rows[0][0] == "MSFT"
        assert data_rows[1][0] == "AAPL"
        # Auto-adjusted widths: each column dimension is set.
        for col_letter, header in zip(("A", "B", "C"), headers):
            width = ws.column_dimensions[col_letter].width
            assert width is not None
            assert width >= len(header)
        wb.close()

    def test_do_get_carries_timeout_call_options(self, tmp_path):
        """The do_get call is bounded: a real FlightCallOptions built
        with ``timeout=120`` is passed, so a wedged-but-connected
        server cannot hang the report (or a chained daily_load_once).
        120 s mirrors the server's own query_timeout_seconds default."""
        config = _report_config(tmp_path)
        table = pa.table(
            {
                "Ticker": ["MSFT"],
                "Avg Trade Money": [1.0],
                "% of Total": [100.0],
            }
        )
        mock_flight = _mock_flight_for_report(table)
        # Wrap the real class so the constructor kwargs are recorded
        # while a genuine FlightCallOptions is still produced (the
        # pyarrow object exposes no readable .timeout attribute).
        mock_flight.FlightCallOptions = MagicMock(
            wraps=flight.FlightCallOptions
        )
        mock_client = mock_flight.FlightClient.return_value

        with patch.object(report_top25_active, "flight", mock_flight):
            code = report_top25_active.run_report(config)

        assert code == 0
        mock_flight.FlightCallOptions.assert_called_once_with(timeout=120)
        mock_client.do_get.assert_called_once()
        options = mock_client.do_get.call_args.kwargs["options"]
        assert isinstance(options, flight.FlightCallOptions)

    def test_flight_down_falls_back_to_direct_readonly_db(self, tmp_path):
        """Server unreachable -> direct read-only DuckDB query.  30
        tickers are seeded so the LIMIT 25 truncation is observable."""
        config = _report_config(tmp_path)
        db_path = config.db_path
        # One row per ticker; OHLC all 1.0 so trade_money == Volume.
        rows = [
            ("2026-09-01", f"T{i:02d}", 1.0, 1.0, 1.0, 1.0, float(1000 - i))
            for i in range(30)
        ]
        _seed_daily_close(db_path, rows)

        mock_flight = _mock_flight_for_report()
        mock_flight.FlightClient.return_value.wait_for_available.side_effect = (
            flight.FlightUnavailableError("connection refused")
        )
        mock_duckdb = MagicMock()
        mock_duckdb.connect = MagicMock(wraps=duckdb.connect)

        with patch.object(
            report_top25_active, "flight", mock_flight
        ), patch.object(report_top25_active, "duckdb", mock_duckdb):
            code = report_top25_active.run_report(config)

        assert code == 0
        # The fallback opened the DB read-only.
        mock_duckdb.connect.assert_called_once_with(db_path, read_only=True)

        out_path = tmp_path / "logs" / "Report_Top_25_Active.xlsx"
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        # LIMIT 25 applied: T00 (vol 1000) .. T24 (vol 976), T25+ dropped.
        assert len(data_rows) == 25
        assert data_rows[0][0] == "T00"
        assert data_rows[0][1] == pytest.approx(1000.0)
        assert data_rows[-1][0] == "T24"
        assert data_rows[-1][1] == pytest.approx(976.0)
        wb.close()

    def test_both_paths_fail_returns_1_no_file(self, tmp_path):
        """Flight unreachable AND the read-only DB connect fails (write
        lock held) -> error logged, exit code 1, no xlsx written."""
        config = _report_config(tmp_path)

        mock_flight = _mock_flight_for_report()
        mock_flight.FlightClient.return_value.wait_for_available.side_effect = (
            flight.FlightUnavailableError("connection refused")
        )
        mock_duckdb = MagicMock()
        mock_duckdb.connect.side_effect = duckdb.IOException(
            "Cannot open file: it is being used by another process"
        )

        with patch.object(
            report_top25_active, "flight", mock_flight
        ), patch.object(report_top25_active, "duckdb", mock_duckdb):
            code = report_top25_active.run_report(config)

        assert code == 1
        assert not (tmp_path / "logs" / "Report_Top_25_Active.xlsx").exists()

    def test_main_returns_1_when_both_paths_fail(self, tmp_path):
        """End-to-end ``main()``: config loads, both query paths fail ->
        process exit code 1."""
        config = _report_config(tmp_path)
        mock_flight = _mock_flight_for_report()
        mock_flight.FlightClient.return_value.wait_for_available.side_effect = (
            flight.FlightUnavailableError("connection refused")
        )
        mock_duckdb = MagicMock()
        mock_duckdb.connect.side_effect = duckdb.IOException("locked")

        with patch.object(
            report_top25_active.PipelineConfig, "from_yaml", return_value=config
        ), patch.object(
            report_top25_active, "_configure_logging"
        ), patch.object(
            report_top25_active, "flight", mock_flight
        ), patch.object(report_top25_active, "duckdb", mock_duckdb):
            assert report_top25_active.main() == 1


class TestDailyLoadOnceReportChaining:
    """``daily_load_once._run_once`` chains the top-25 report after a
    successful load — via either path — and a report failure never
    changes the load's exit code."""

    @pytest.fixture
    def config(self, tmp_path):
        return PipelineConfig(
            database={
                "duckdb_dir": str(tmp_path),
                "duckdb_filename": "test.duckdb",
            },
            scheduler=SchedulerConfig(enabled=True),
        )

    @pytest.mark.parametrize("status", ["SUCCESS", "PARTIAL"])
    def test_flight_success_chains_report(self, config, status):
        mock_flight = _mock_flight_module()
        mock_flight.FlightClient.return_value.do_action.return_value = iter(
            [_flight_result({"status": status, "rows_upserted": 5})]
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ), patch.object(
            daily_load_once, "report_top25_active"
        ) as mock_report:
            mock_report.run_report.return_value = 0
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_report.run_report.assert_called_once_with(config)

    def test_flight_failure_does_not_chain_report(self, config):
        """Load FAILED via the service -> no report run at all."""
        mock_flight = _mock_flight_module()
        mock_flight.FlightClient.return_value.do_action.side_effect = (
            flight.FlightServerError("server exploded")
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ), patch.object(
            daily_load_once, "report_top25_active"
        ) as mock_report:
            code = daily_load_once._run_once(config)

        assert code == 1
        mock_report.run_report.assert_not_called()

    @pytest.mark.parametrize(
        "report_outcome",
        [1, RuntimeError("xlsx write exploded")],
        ids=["returns_1", "raises"],
    )
    def test_report_failure_keeps_load_exit_0(self, config, report_outcome):
        """A report failure (non-zero return OR exception) is logged but
        must not change the successful load's exit code."""
        mock_flight = _mock_flight_module()
        mock_flight.FlightClient.return_value.do_action.return_value = iter(
            [_flight_result({"status": "SUCCESS", "rows_upserted": 5})]
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ), patch.object(
            daily_load_once, "report_top25_active"
        ) as mock_report:
            if isinstance(report_outcome, Exception):
                mock_report.run_report.side_effect = report_outcome
            else:
                mock_report.run_report.return_value = report_outcome
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_report.run_report.assert_called_once_with(config)

    def test_direct_fallback_success_chains_report(self, config):
        """Flight unreachable -> direct load path succeeds -> report
        chained with the same config."""
        mock_flight = _mock_flight_module()
        mock_flight.FlightClient.return_value.wait_for_available.side_effect = (
            flight.FlightUnavailableError("connection refused")
        )

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "PipelineScheduler"
        ) as mock_sched_cls, patch.object(
            daily_load_once, "report_top25_active"
        ) as mock_report:
            mock_report.run_report.return_value = 0
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_sched = mock_sched_cls.return_value
        mock_sched._execute_daily_load.assert_called_once_with(
            mock_sched._create_db_manager.return_value, force=True
        )
        mock_report.run_report.assert_called_once_with(config)

    def test_scheduler_disabled_does_not_chain_report(self, tmp_path):
        """enabled=false -> exit 0 but no load ran, so no report."""
        config = PipelineConfig(
            database={
                "duckdb_dir": str(tmp_path),
                "duckdb_filename": "test.duckdb",
            },
            scheduler=SchedulerConfig(enabled=False),
        )
        mock_flight = _mock_flight_module()

        with patch.object(daily_load_once, "flight", mock_flight), patch.object(
            daily_load_once, "report_top25_active"
        ) as mock_report:
            code = daily_load_once._run_once(config)

        assert code == 0
        mock_report.run_report.assert_not_called()
